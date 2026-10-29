from io import BytesIO
from pathlib import Path
import fitz
from docx import Document

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from a text-based PDF or DOCX resume."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        chunks = []
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            for page in pdf:
                chunks.append(page.get_text("text"))
        return "\n".join(chunks).strip()
    if suffix == ".docx":
        doc = Document(BytesIO(file_bytes))
        chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(chunks).strip()
    raise ValueError("Please upload a PDF or DOCX file.")
